"""Rebuild the two demo workbooks in dist/demo from the Case 1 template.

Usage: python3 -E -P scripts/demo-workbooks/build.py <template.xlsx> dist/demo
Rows 1-3 of each sheet (titles, headers, the grey row-3 example) and the other sheets,
including the fixed Existing stores list, are kept exactly as in the template.
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data  # noqa: E402


def fill(template, path_out, territories, apps):
    wb = openpyxl.load_workbook(template)
    for sheet, rows in (('Applications', apps), ('Territories', territories)):
        ws = wb[sheet]
        if ws.max_row >= 4:
            ws.delete_rows(4, ws.max_row - 3)
        for r, row in enumerate(rows, start=4):
            for c, value in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=value)
    wb.save(path_out)
    print(path_out, len(apps), 'applicants', len(territories), 'territories')


if __name__ == '__main__':
    template, out_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)
    fill(template, os.path.join(out_dir, 'demo-1-clean.xlsx'), data.DEMO1_TERRITORIES, data.DEMO1_APPS)
    fill(template, os.path.join(out_dir, 'demo-2-messy.xlsx'), data.DEMO2_TERRITORIES, data.DEMO2_APPS)
